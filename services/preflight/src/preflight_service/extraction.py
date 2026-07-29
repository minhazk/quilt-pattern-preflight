from __future__ import annotations

import io
import mimetypes
import re
from pathlib import Path

import pdfplumber
from docx import Document

from .measurements import normalize_fraction_text
from .models import BoundingBox, ExtractedEntity, ExtractionResult, SourceReference

MAX_UPLOAD_BYTES = 15 * 1024 * 1024
SUPPORTED = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pdf": "application/pdf",
}
DIMENSION_PATTERN = re.compile(
    r"(?P<width>\d+(?:\s+\d+/\d+|[¼½¾⅛⅜⅝⅞]|\.\d+)?)\s*"
    r"(?:in(?:ches?)?|[\"″])?\s*(?:x|×|by)\s*"
    r"(?P<height>WOF|\d+(?:\s+\d+/\d+|[¼½¾⅛⅜⅝⅞]|\.\d+)?)",
    re.IGNORECASE,
)
QUANTITY_PATTERN = re.compile(
    r"(?:cut|make|quantity|qty)[^\n]{0,28}?\(?(\d{1,4})\)?",
    re.IGNORECASE,
)


class ExtractionError(ValueError):
    """Safe, customer-readable extraction failure."""


def validate_upload(filename: str, content_type: str | None, content: bytes) -> str:
    if len(content) > MAX_UPLOAD_BYTES:
        raise ExtractionError("The file is larger than the 15 MB pilot limit.")
    if not content:
        raise ExtractionError("The uploaded file is empty.")
    safe_name = Path(filename).name
    if safe_name != filename or "\x00" in filename:
        raise ExtractionError("The filename contains unsupported path characters.")
    extension = Path(safe_name).suffix.casefold()
    if extension not in SUPPORTED:
        raise ExtractionError("Upload a DOCX or text-layer PDF file.")
    expected_mime = SUPPORTED[extension]
    guessed_mime, _ = mimetypes.guess_type(safe_name)
    accepted_mimes = {expected_mime, guessed_mime, "application/octet-stream", None}
    if content_type not in accepted_mimes:
        raise ExtractionError("The file type does not match its extension.")
    if extension == ".pdf" and not content.startswith(b"%PDF"):
        raise ExtractionError("The file does not contain a valid PDF header.")
    if extension == ".docx" and not content.startswith(b"PK"):
        raise ExtractionError("The file does not contain a valid DOCX package.")
    return extension[1:]


def _entities_from_text(
    text: str,
    *,
    page: int | None,
    section: str | None,
    method: str,
    bounding_box: BoundingBox | None = None,
) -> list[ExtractedEntity]:
    entities: list[ExtractedEntity] = []
    for match in DIMENSION_PATTERN.finditer(text):
        width = normalize_fraction_text(match.group("width"))
        raw_height = match.group("height")
        height = "WOF" if raw_height.casefold() == "wof" else normalize_fraction_text(raw_height)
        excerpt = text[max(0, match.start() - 45) : match.end() + 45].strip()
        entities.append(
            ExtractedEntity(
                entity_type="dimension",
                value=match.group(0),
                normalized_value=f"{width} x {height}",
                source=SourceReference(
                    page=page,
                    section=section,
                    excerpt=excerpt,
                    bounding_box=bounding_box,
                ),
                extraction_method=method,
                confidence=0.9 if section else 0.82,
            )
        )
    for match in QUANTITY_PATTERN.finditer(text):
        excerpt = text[max(0, match.start() - 35) : match.end() + 50].strip()
        entities.append(
            ExtractedEntity(
                entity_type="quantity",
                value=match.group(1),
                normalized_value=match.group(1),
                source=SourceReference(
                    page=page,
                    section=section,
                    excerpt=excerpt,
                    bounding_box=bounding_box,
                ),
                extraction_method=method,
                confidence=0.83 if section else 0.75,
            )
        )
    return entities


def extract_docx(filename: str, content: bytes) -> ExtractionResult:
    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise ExtractionError(
            "The DOCX could not be opened. Export a new copy and try again."
        ) from exc
    entities: list[ExtractedEntity] = []
    text_parts: list[str] = []
    current_heading: str | None = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        text_parts.append(text)
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            current_heading = text
        entities.extend(
            _entities_from_text(
                text,
                page=None,
                section=current_heading,
                method="docx_paragraph_regex_v1",
            )
        )
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if not row_text.strip(" |"):
                continue
            text_parts.append(row_text)
            entities.extend(
                _entities_from_text(
                    row_text,
                    page=None,
                    section=f"Table {table_index}, row {row_index}",
                    method="docx_table_regex_v1",
                )
            )
    plain_text = "\n".join(text_parts)
    if len(plain_text.strip()) < 40:
        raise ExtractionError("The DOCX contains too little extractable text for preflight.")
    return ExtractionResult(
        filename=Path(filename).name,
        file_type="docx",
        page_count=max(1, len(document.sections)),
        text_layer_available=True,
        structural_suitability="suitable" if entities else "review",
        plain_text_length=len(plain_text),
        entities=entities,
        warnings=(
            [] if entities else ["No measurements or quantities were confidently extracted."]
        ),
    )


def extract_pdf(filename: str, content: bytes) -> ExtractionResult:
    try:
        pdf = pdfplumber.open(io.BytesIO(content))
    except Exception as exc:
        raise ExtractionError(
            "The PDF could not be opened. It may be corrupted or encrypted."
        ) from exc
    entities: list[ExtractedEntity] = []
    text_parts: list[str] = []
    with pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            text_parts.append(text)
            if text:
                bbox = BoundingBox(
                    x=0,
                    y=0,
                    width=float(page.width),
                    height=float(page.height),
                )
                entities.extend(
                    _entities_from_text(
                        text,
                        page=page_number,
                        section=None,
                        method="pdf_text_regex_v1",
                        bounding_box=bbox,
                    )
                )
        page_count = len(pdf.pages)
    plain_text = "\n".join(text_parts).strip()
    if len(plain_text) < 40:
        raise ExtractionError(
            "No usable text layer was found. Scanned and image-only PDFs are "
            "outside the current preflight scope."
        )
    return ExtractionResult(
        filename=Path(filename).name,
        file_type="pdf",
        page_count=page_count,
        text_layer_available=True,
        structural_suitability="suitable" if entities else "review",
        plain_text_length=len(plain_text),
        entities=entities,
        warnings=(
            [] if entities else ["No measurements or quantities were confidently extracted."]
        ),
    )


def extract_document(filename: str, content_type: str | None, content: bytes) -> ExtractionResult:
    file_type = validate_upload(filename, content_type, content)
    if file_type == "docx":
        return extract_docx(filename, content)
    return extract_pdf(filename, content)
